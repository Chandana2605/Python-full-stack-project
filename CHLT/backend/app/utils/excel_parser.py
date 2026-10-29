from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from ..exceptions.http_exceptions import bad_request_exception


def parse_excel_file(file_bytes: bytes, required_columns: list[str]) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(filename=BytesIO(file_bytes), data_only=True)
        worksheet = workbook.active
    except Exception as exc:
        raise bad_request_exception(f"Invalid Excel file: {exc}") from exc

    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise bad_request_exception("Excel file is empty")

    header_row = rows[0]
    if not header_row:
        raise bad_request_exception("Excel header row is missing")

    normalized_headers = [
        str(column).strip().lower().replace(" ", "_") if column is not None else ""
        for column in header_row
    ]
    missing_columns = [column for column in required_columns if column not in normalized_headers]
    if missing_columns:
        raise bad_request_exception(f"Missing required columns: {', '.join(missing_columns)}")

    records: list[dict[str, Any]] = []
    for row in rows[1:]:
        if row is None or all(cell is None for cell in row):
            continue
        record = {}
        for index, header in enumerate(normalized_headers):
            if not header:
                continue
            value = row[index] if index < len(row) else None
            record[header] = value
        records.append(record)

    return records
