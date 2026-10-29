import json
from io import BytesIO
from typing import Any

import asyncpg
from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook

from app.db.postgresql import get_database
from app.schemas.common import SimpleResolveRequest, utc_now
from app.utils.ids import new_id

review_cache: list[dict[str, Any]] = []


def _normalize_header(value: Any) -> str:
    return str(value).strip().lower() if value is not None else ""


def _normalize_key(value: str) -> str:
    return value.strip().lower().replace("-", "_").replace(" ", "_")


def _compact_key(value: str) -> str:
    return "".join(char for char in value.strip().lower() if char.isalnum())


def _normalize_string(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _normalize_email(value: Any) -> str:
    return str(value).strip().lower() if value is not None else ""


def _parse_attended(value: Any) -> bool:
    normalized = _normalize_string(value).lower()
    if normalized == "":
        return True
    return normalized in {"yes", "true", "1", "y"}


def _first_value(data: dict[str, Any], *keys: str) -> Any:
    normalized_data = {_normalize_key(key): value for key, value in data.items()}
    compact_data = {_compact_key(key): value for key, value in data.items()}
    for key in keys:
        value = normalized_data.get(_normalize_key(key))
        if value is None:
            value = compact_data.get(_compact_key(key))
        if value is not None and _normalize_string(value):
            return value
    return None


def _extract_candidate_id(data: dict[str, Any]) -> str:
    return _normalize_string(
        _first_value(data, "candidate_id", "candidate id", "candidateid", "id", "employee_id", "employee id")
    )


def _extract_email(data: dict[str, Any]) -> str:
    return _normalize_email(
        _first_value(
            data,
            "candidate_email",
            "candidate email",
            "candidateemail",
            "student_email",
            "student email",
            "studentemail",
            "email",
            "email_id",
            "email id",
            "emailid",
            "email_address",
            "email address",
            "emailaddress",
            "registered_email",
            "registered email",
            "registeredemail",
            "personal_email",
            "personal email",
            "personalemail",
        )
    )


def _extract_name(data: dict[str, Any]) -> str:
    return _normalize_string(
        _first_value(data, "name", "candidate_name", "candidate name", "student_name", "student name", "full_name", "full name")
    )


def _extract_college(data: dict[str, Any]) -> str:
    return _normalize_string(
        _first_value(data, "college", "college_name", "college name", "institution", "institution_name", "institution name")
    )


def _extract_float(data: dict[str, Any], default: float, *keys: str) -> float:
    value = _first_value(data, *keys)
    if value in (None, ""):
        return default
    return float(value)


async def _find_candidate_by_email(db, email: str) -> dict[str, Any] | None:
    if not email:
        return None
    return await db.candidates.find_one(
        {"$or": [{"email": email}, {"alternate_emails": email}]}
    )


async def _find_candidate_by_name_college(db, name: str, college: str) -> dict[str, Any] | None:
    if not name:
        return None

    if name and college:
        possible = await db.candidates.find(
            {
                "name": {"$regex": name, "$options": "i"},
                "college": {"$regex": college, "$options": "i"},
            }
        ).to_list(length=5)
        if len(possible) == 1:
            return possible[0]

    possible = await db.candidates.find(
        {"name": {"$regex": name, "$options": "i"}}
    ).to_list(length=5)
    if len(possible) == 1:
        return possible[0]

    return None


async def _collect_possible_matches(db, candidate_id: str, email: str, name: str, college: str) -> list[dict[str, str]]:
    matches: list[dict[str, str]] = []
    seen_ids: set[str] = set()

    if candidate_id:
        candidate = await db.candidates.find_one({"id": candidate_id})
        if candidate:
            seen_ids.add(candidate["id"])
            matches.append({"id": candidate["id"], "name": candidate.get("name", ""), "email": candidate.get("email", "")})

    if email:
        candidates = await db.candidates.find(
            {"$or": [{"email": email}, {"alternate_emails": email}]}
        ).to_list(length=5)
        for candidate in candidates:
            if candidate["id"] not in seen_ids:
                seen_ids.add(candidate["id"])
                matches.append({"id": candidate["id"], "name": candidate.get("name", ""), "email": candidate.get("email", "")})

    if name:
        query: dict[str, Any] = {"name": {"$regex": name, "$options": "i"}}
        if college:
            query["college"] = {"$regex": college, "$options": "i"}
        possible = await db.candidates.find(query).to_list(length=5)
        for candidate in possible:
            if candidate["id"] not in seen_ids:
                seen_ids.add(candidate["id"])
                matches.append({"id": candidate["id"], "name": candidate.get("name", ""), "email": candidate.get("email", "")})

    return matches[:5]


async def _save_baseline_for_candidate(
    db,
    candidate: dict[str, Any],
    data: dict[str, Any],
    row_email: str,
    mapped_email: str | None,
    is_attended: bool,
    score: float,
    user_id: str,
) -> None:
    alt = candidate.get("alternate_emails") or []
    if row_email and row_email != candidate["email"] and row_email not in alt:
        alt.append(row_email)
    if mapped_email:
        normalized_mapped = _normalize_email(mapped_email)
        if normalized_mapped and normalized_mapped != candidate["email"] and normalized_mapped not in alt:
            alt.append(normalized_mapped)
        mapped_email = normalized_mapped

    baseline_email_map = candidate.get("baseline_email_map") or {}
    if row_email and row_email != candidate["email"]:
        baseline_email_map[row_email] = candidate["email"]
    if mapped_email and mapped_email != candidate["email"]:
        baseline_email_map[mapped_email] = candidate["email"]

    await db.update_candidate_row(
        {"id": candidate["id"]},
        {
            "alternate_emails": json.dumps(alt),
            "baseline_data": json.dumps({}),
            "baseline_email_map": json.dumps(baseline_email_map),
            "pre_onboarding_score": score,
            "updated_at": utc_now(),
        },
    )

    max_score = _extract_float(data, 100, "max_score", "max score", "maximum_score", "maximum score", "total_score", "total score")
    pass_percentage = _extract_float(
        data, 0, "pass_percentage", "pass percentage", "passing_percentage", "passing percentage", "cutoff", "cut_off", "cut off"
    )
    passed = is_attended and score >= pass_percentage

    result = {
        "id": new_id(),
        "candidate_id": candidate["id"],
        "batch_id": candidate.get("batch_id"),
        "phase": "PRE_ONBOARDING",
        "exam_type": "BASELINE",
        "exam_name": "BASELINE",
        "attempt": 1,
        "score": score,
        "max_score": max_score,
        "pass_percentage": pass_percentage,
        "passed": passed,
        "candidate_status_after": candidate.get("status") or "ELIGIBLE",
        "timeline_badge": None,
        "created_by": user_id,
        "created_at": utc_now(),
    }
    try:
        await db.exam_results.insert_one(result)
    except asyncpg.IntegrityConstraintViolationError:
        pass


async def upload_baseline(file: UploadFile, user_id: str) -> dict[str, Any]:
    if not file.filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Only .xlsx supported")

    workbook = load_workbook(BytesIO(await file.read()))
    sheet = workbook.active

    headers = [_normalize_header(cell.value) for cell in sheet[1]]
    db = get_database()
    rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]

    processed = 0
    alternate_added = 0
    review_required: list[dict[str, Any]] = []

    for idx, data in enumerate(rows, start=2):
        candidate_id = _extract_candidate_id(data)
        email = _extract_email(data)
        name = _extract_name(data)
        college = _extract_college(data)

        if not candidate_id and not email and not name:
            review_required.append(
                {
                    "row": idx,
                    "email": email,
                    "name": name,
                    "college": college,
                    "data": data,
                    "possible_matches": [],
                    "error": "Missing candidate_id, candidate_email, email, or name",
                }
            )
            continue

        candidate = None
        if candidate_id:
            candidate = await db.candidates.find_one({"id": candidate_id})
        if not candidate and email:
            candidate = await _find_candidate_by_email(db, email)
        if not candidate:
            candidate = await _find_candidate_by_name_college(db, name, college)

        if not candidate:
            review_required.append(
                {
                    "row": idx,
                    "email": email,
                    "name": name,
                    "college": college,
                    "data": data,
                    "possible_matches": await _collect_possible_matches(db, candidate_id, email, name, college),
                }
            )
            continue

        alt = candidate.get("alternate_emails") or []
        if email and email != candidate["email"] and email not in alt:
            alt.append(email)
            await db.candidates.update_one(
                {"id": candidate["id"]},
                {"$set": {"alternate_emails": alt}},
            )
            alternate_added += 1

        is_attended = _parse_attended(data.get("attended"))
        score = _extract_float(data, 0, "score", "marks", "obtained_score", "obtained score", "percentage")
        if not is_attended:
            score = 0.0

        await _save_baseline_for_candidate(db, candidate, data, email, None, is_attended, score, user_id)
        processed += 1

    review_cache.clear()
    review_cache.extend(review_required)

    return {
        "processed": processed,
        "alternate_emails_added": alternate_added,
        "review_required": review_required[:50],
    }


async def _find_candidate_for_review(db, email: str, name: str, college: str) -> dict[str, Any] | None:
    candidate = await _find_candidate_by_email(db, email) if email else None
    if candidate:
        return candidate
    return await _find_candidate_by_name_college(db, name, college)


async def _find_only_candidate_without_baseline(db) -> dict[str, Any] | None:
    candidates = await db.candidates.find({}).to_list(length=10000)
    candidates_without_baseline = [
        candidate
        for candidate in candidates
        if not (candidate.get("baseline_data") or {}).get("baseline_exam")
    ]
    if len(candidates_without_baseline) == 1:
        return candidates_without_baseline[0]
    return None


async def _resolve_review_item(
    db,
    item: dict[str, Any],
    mapping: Any | None,
    user_id: str,
    auto_map: bool = False,
) -> bool:
    data = item.get("data") or {}
    email = item.get("email") or _extract_email(data)
    name = item.get("name") or _extract_name(data)
    college = item.get("college") or _extract_college(data)
    mapped_email = _normalize_email(mapping.mapped_email) if mapping and mapping.mapped_email else None
    candidate_id = mapping.candidate_id if mapping else None
    possible_matches = item.get("possible_matches", [])

    if auto_map and not possible_matches:
        possible_matches = await _collect_possible_matches(db, _extract_candidate_id(data), email, name, college)

    if mapping and mapping.selected_index is not None:
        if 0 <= mapping.selected_index < len(possible_matches):
            candidate_id = possible_matches[mapping.selected_index]["id"]

    if auto_map and not candidate_id and possible_matches:
        candidate_id = possible_matches[0]["id"]

    if not candidate_id and len(possible_matches) == 1:
        candidate_id = possible_matches[0]["id"]

    candidate = None
    if candidate_id:
        candidate = await db.candidates.find_one({"id": candidate_id})
    if not candidate and mapped_email:
        candidate = await _find_candidate_by_email(db, _normalize_email(mapped_email))
    if not candidate:
        candidate = await _find_candidate_for_review(db, email, name, college)
    if auto_map and not candidate:
        candidate = await _find_only_candidate_without_baseline(db)
    if not candidate:
        return False

    is_attended = _parse_attended(data.get("attended"))
    score = _extract_float(data, 0, "score", "marks", "obtained_score", "obtained score", "percentage")
    if not is_attended:
        score = 0.0

    await _save_baseline_for_candidate(db, candidate, data, email, mapped_email, is_attended, score, user_id)
    return True


def _review_summary(item: dict[str, Any]) -> dict[str, Any]:
    data = item.get("data") or {}
    return {
        "row": item.get("row"),
        "email": item.get("email") or _extract_email(data),
        "name": item.get("name") or _extract_name(data),
        "college": item.get("college") or _extract_college(data),
        "possible_matches": item.get("possible_matches", []),
        "error": item.get("error"),
    }


async def resolve_mismatches(payload: SimpleResolveRequest, user_id: str) -> dict[str, Any]:
    db = get_database()
    row_mapping_by_row = {mapping.row: mapping for mapping in payload.row_mappings}

    if payload.resolve_all:
        items_to_resolve = list(review_cache)
    else:
        rows_to_resolve = set(payload.rows) | set(row_mapping_by_row.keys())
        items_to_resolve = [item for item in review_cache if item.get("row") in rows_to_resolve]

    resolved = 0
    resolved_rows: list[int] = []

    errors: list[dict[str, Any]] = []

    for item in items_to_resolve:
        mapping = row_mapping_by_row.get(item.get("row"))
        try:
            if await _resolve_review_item(db, item, mapping, user_id, auto_map=payload.resolve_all):
                resolved += 1
                resolved_rows.append(item.get("row"))
        except Exception as exc:
            item["error"] = str(exc)
            errors.append(_review_summary(item))

    if resolved_rows:
        review_cache[:] = [item for item in review_cache if item.get("row") not in resolved_rows]

    return {
        "resolved": resolved,
        "remaining": len(review_cache),
        "unresolved": [_review_summary(item) for item in review_cache[:50]],
        "errors": errors[:20],
    }
