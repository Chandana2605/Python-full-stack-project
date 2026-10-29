import json
from io import BytesIO
from typing import Any

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook

from app.db.postgresql import get_database
from app.schemas.common import utc_now
from app.utils.ids import new_id


def _to_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    return float(value)


def _to_int(value: Any) -> int | None:
    if value in (None, ""):
        return None
    return int(value)


def _to_bool(value: Any) -> bool | None:
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {"true", "yes", "y", "1"}:
        return True
    if normalized in {"false", "no", "n", "0"}:
        return False
    raise ValueError(f"Invalid boolean value: {value}")


def _to_string(value: Any) -> str | None:
    if value in (None, ""):
        return None
    return str(value).strip()


def _normalize_key(value: str) -> str:
    return value.strip().lower().replace("-", "_").replace(" ", "_")


def _compact_key(value: str) -> str:
    return "".join(char for char in value.strip().lower() if char.isalnum())


def _first_value(data: dict[str, Any], *keys: str) -> Any:
    normalized_data = {_normalize_key(key): value for key, value in data.items()}
    compact_data = {_compact_key(key): value for key, value in data.items()}
    for key in keys:
        value = normalized_data.get(_normalize_key(key))
        if value is None:
            value = compact_data.get(_compact_key(key))
        if value not in (None, ""):
            return value
    return None


def _structured_candidate_fields(data: dict[str, Any], now: Any) -> dict[str, Any]:
    email = str(
        _first_value(
            data,
            "email",
            "email_id",
            "email id",
            "email_address",
            "email address",
            "candidate_email",
            "candidate email",
        )
        or ""
    ).strip().lower()
    name = _to_string(_first_value(data, "name", "candidate_name", "candidate name", "student_name", "student name"))
    college = _to_string(_first_value(data, "college", "college_name", "college name", "institution", "institution_name"))
    percentage = _to_float(_first_value(data, "percentage", "marks", "score"))
    eligible_phase = str(_first_value(data, "eligible_phase", "eligible phase") or "PRE_ONBOARDING").strip().upper()
    intake_year = _to_int(_first_value(data, "intake_year", "intake year"))
    graduation_year = _to_int(_first_value(data, "graduation_year", "graduation year"))
    tech_stack = _to_string(_first_value(data, "tech_stack", "tech stack", "technology", "technology_stack", "technology stack"))
    source = _to_string(_first_value(data, "source"))
    offer_accepted = _to_bool(_first_value(data, "offer_accepted", "offer accepted"))

    if not email:
        raise ValueError("Email is required")

    return {
        "email": email,
        "name": name,
        "college": college,
        "percentage": percentage,
        "eligible_phase": eligible_phase,
        "intake_year": intake_year,
        "graduation_year": graduation_year,
        "tech_stack": tech_stack,
        "source": source,
        "offer_accepted": offer_accepted,
        "updated_at": now,
    }


async def upload_candidates(file: UploadFile) -> dict:
    if not file.filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Only .xlsx supported")

    workbook = load_workbook(BytesIO(await file.read()))
    sheet = workbook.active

    headers = [str(cell.value).strip().lower() if cell.value else "" for cell in sheet[1]]

    db = get_database()
    rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    inserted = 0
    updated = 0
    errors = []

    for idx, data in enumerate(rows, start=2):

        try:
            now = utc_now()
            candidate_fields = _structured_candidate_fields(data, now)
            email = candidate_fields["email"]
            existing = await db.candidates.find_one({"email": email})

            if existing:
                updates = {
                    key: value
                    for key, value in candidate_fields.items()
                    if key == "updated_at" or value is not None
                }
                await db.update_candidate_row({"email": email}, updates)
                updated += 1
                continue

            candidate = {
                "id": new_id(),
                "email": email,
                "name": candidate_fields["name"] or email.split("@")[0],
                "eligible_phase": candidate_fields["eligible_phase"],
                "intake_year": candidate_fields["intake_year"],
                "college": candidate_fields["college"],
                "graduation_year": candidate_fields["graduation_year"],
                "tech_stack": candidate_fields["tech_stack"],
                "percentage": candidate_fields["percentage"],
                "batch_id": None,
                "uid": None,
                "employee_email": None,
                "offer_accepted": candidate_fields["offer_accepted"] if candidate_fields["offer_accepted"] is not None else False,
                "source": candidate_fields["source"],
                "status": "ELIGIBLE",
                "pre_onboarding_score": None,
                "post_onboarding_score": None,
                "knowledge_check_python_score": None,
                "knowledge_check_java_score": None,
                "marks": {},
                "alternate_emails": [],
                "baseline_data": {},
                "baseline_email_map": {},
                "created_at": now,
                "updated_at": now,
            }

            await db.insert_candidate_row(candidate)
            inserted += 1

        except Exception as exc:
            errors.append({"row": idx, "error": str(exc)})

    return {
        "inserted": inserted,
        "updated": updated,
        "errors": errors[:20],
    }


async def get_candidates(
    search: str = None,
    min_score: float = None,
    max_score: float = None,
    attended: bool = None,
    page: int = 1,
    limit: int = 10,
):
    db = get_database()

    conditions = []
    values = []
    idx = 1

    if search:
        conditions.append(f"(LOWER(name) LIKE ${idx} OR LOWER(email) LIKE ${idx})")
        values.append(f"%{search.lower()}%")
        idx += 1

    if min_score is not None:
        conditions.append(f"pre_onboarding_score >= ${idx}")
        values.append(min_score)
        idx += 1

    if max_score is not None:
        conditions.append(f"pre_onboarding_score <= ${idx}")
        values.append(max_score)
        idx += 1

    if attended is not None:
        if attended:
            conditions.append("pre_onboarding_score > 0")
        else:
            conditions.append("(pre_onboarding_score IS NULL OR pre_onboarding_score = 0)")

    where_clause = ""
    if conditions:
        where_clause = "WHERE " + " AND ".join(conditions)

    page = max(page, 1)
    limit = min(max(limit, 1), 100)
    offset = (page - 1) * limit
    limit_param = idx
    offset_param = idx + 1

    query = f"""
        SELECT id, name, email, college, pre_onboarding_score, alternate_emails
        FROM candidates
        {where_clause}
        ORDER BY created_at DESC
        LIMIT ${limit_param} OFFSET ${offset_param}
    """

    count_query = f"""
        SELECT COUNT(*) FROM candidates
        {where_clause}
    """

    rows = await db.fetch(query, *values, limit, offset)
    total = await db.fetchval(count_query, *values)

    items = []
    for row in rows:
        score = row["pre_onboarding_score"]
        alternate_emails = row["alternate_emails"] or []
        if isinstance(alternate_emails, str):
            try:
                alternate_emails = json.loads(alternate_emails)
            except json.JSONDecodeError:
                alternate_emails = []

        items.append(
            {
                "id": row["id"],
                "name": row["name"],
                "email": row["email"],
                "college": row["college"],
                "alternate_emails": alternate_emails,
                "score": score,
                "attended": score is not None and score > 0,
            }
        )

    return {
        "total": total,
        "page": page,
        "limit": limit,
        "items": items,
    }
