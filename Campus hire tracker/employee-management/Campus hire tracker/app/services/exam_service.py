from io import BytesIO
from io import StringIO

import asyncpg
from fastapi import HTTPException, UploadFile, status
from openpyxl import load_workbook

from app.core.enums import CandidateStatus, ExamPhase, ExamType, TimelineEventType
from app.db.postgresql import get_database
from app.schemas.batches import ExamScoreInput
from app.schemas.common import utc_now
from app.services.batch_service import assert_batch_owner
from app.services.timeline_service import add_timeline_event
from app.utils.ids import new_id


def _exam_name(payload: ExamScoreInput) -> str:
    return payload.custom_exam_name or payload.exam_type.value


def _timeline_badge(phase: ExamPhase, exam_type: ExamType, candidate_status: CandidateStatus) -> str | None:
    if phase == ExamPhase.PRE_ONBOARDING:
        if candidate_status == CandidateStatus.ELIGIBLE_FOR_OFFER:
            return "Eligible for Offer"
        if candidate_status == CandidateStatus.RECOMMENDED_FOR_LOI_CANCELLATION:
            return "Recommended for LOI Cancellation"
        return None
    if exam_type == ExamType.FINAL_EXAM:
        if candidate_status == CandidateStatus.DEPLOYMENT_READY:
            return "Deployment Ready"
        if candidate_status == CandidateStatus.REFERRED_TO_HR:
            return "Referred to HR"
    return None


def _status_after_exam(phase: ExamPhase, attempt: int, passed: bool) -> CandidateStatus:
    if phase == ExamPhase.PRE_ONBOARDING:
        if passed:
            return CandidateStatus.ELIGIBLE_FOR_OFFER
        return (
            CandidateStatus.PRE_ONBOARDING_REATTEMPT_REQUIRED
            if attempt == 1
            else CandidateStatus.RECOMMENDED_FOR_LOI_CANCELLATION
        )
    if passed:
        return CandidateStatus.DEPLOYMENT_READY
    return CandidateStatus.POST_ONBOARDING_REATTEMPT_REQUIRED if attempt == 1 else CandidateStatus.REFERRED_TO_HR


async def _validate_attempt(candidate_id: str, phase: ExamPhase, exam_name: str, attempt: int) -> None:
    db = get_database()
    existing = await db.exam_results.find(
        {"candidate_id": candidate_id, "phase": phase.value, "exam_name": exam_name}
    ).to_list(length=10)
    if any(result["passed"] for result in existing):
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Reattempt is disabled after pass")
    if attempt == 2:
        first = next((result for result in existing if result["attempt"] == 1), None)
        if not first:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="First attempt is required first")
        if first["passed"]:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Candidate already passed")
    if attempt == 1 and existing:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="First attempt already recorded")


async def record_exam_score(batch_id: str, owner_id: str, payload: ExamScoreInput) -> dict:
    batch = await assert_batch_owner(batch_id, owner_id)
    db = get_database()
    candidate = await db.candidates.find_one({"id": payload.candidate_id, "batch_id": batch_id})
    if not candidate:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Candidate not found in this batch")
    if batch["phase"] != payload.phase.value:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Exam phase must match batch phase")
    if candidate["eligible_phase"] != payload.phase.value:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Candidate is not eligible for this phase")
    if payload.phase == ExamPhase.POST_ONBOARDING and (not candidate.get("uid") or not candidate.get("employee_email")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="UID and employee email must be updated before post-onboarding exams",
        )
    exam_name = _exam_name(payload)
    await _validate_attempt(payload.candidate_id, payload.phase, exam_name, payload.attempt)
    percentage = (payload.score / payload.max_score) * 100
    passed = percentage >= payload.pass_percentage
    candidate_status = _status_after_exam(payload.phase, payload.attempt, passed)
    timeline_badge = _timeline_badge(payload.phase, payload.exam_type, candidate_status)
    result = {
        "id": new_id(),
        "candidate_id": payload.candidate_id,
        "batch_id": batch_id,
        "phase": payload.phase.value,
        "exam_type": payload.exam_type.value,
        "exam_name": exam_name,
        "attempt": payload.attempt,
        "score": payload.score,
        "max_score": payload.max_score,
        "pass_percentage": payload.pass_percentage,
        "passed": passed,
        "candidate_status_after": candidate_status.value,
        "timeline_badge": timeline_badge,
        "created_by": owner_id,
        "created_at": utc_now(),
    }
    try:
        await db.exam_results.insert_one(result)
    except asyncpg.IntegrityConstraintViolationError as exc:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Exam attempt already exists") from exc
    score_field = (
        "pre_onboarding_score" if payload.phase == ExamPhase.PRE_ONBOARDING else "post_onboarding_score"
    )
    await db.candidates.update_one(
        {"id": payload.candidate_id},
        {"$set": {score_field: payload.score, "status": candidate_status.value, "updated_at": utc_now()}},
    )
    await add_timeline_event(
        payload.candidate_id,
        owner_id,
        TimelineEventType.EXAM_UPDATED,
        f"{payload.phase.value} exam {exam_name} attempt {payload.attempt} recorded",
        {
            "exam_type": payload.exam_type.value,
            "score": payload.score,
            "max_score": payload.max_score,
            "passed": passed,
            "candidate_status_after": candidate_status.value,
            "timeline_badge": timeline_badge,
        },
    )
    return result


async def upload_exam_scores(batch_id: str, owner_id: str, file: UploadFile) -> list[dict]:
    content = await file.read()
    workbook = load_workbook(filename=BytesIO(content))
    sheet = workbook.active
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
    required = ["candidate_id", "phase", "exam_type", "attempt", "score", "max_score"]
    missing = [header for header in required if header not in headers]
    if missing:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Missing columns: {missing}")
    results: list[dict] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        data = dict(zip(headers, row, strict=False))
        if not data.get("candidate_id"):
            continue
        payload = ExamScoreInput(
            candidate_id=str(data["candidate_id"]),
            phase=ExamPhase(str(data["phase"])),
            exam_type=ExamType(str(data["exam_type"])),
            custom_exam_name=str(data["custom_exam_name"]) if data.get("custom_exam_name") else None,
            attempt=int(data["attempt"]),
            score=float(data["score"]),
            max_score=float(data["max_score"]),
            pass_percentage=float(data.get("pass_percentage") or 60),
        )
        results.append(await record_exam_score(batch_id, owner_id, payload))
    return results


async def list_exam_results(batch_id: str, owner_id: str | None = None) -> list[dict]:
    if owner_id:
        await assert_batch_owner(batch_id, owner_id)
    return await get_database().exam_results.find({"batch_id": batch_id}).to_list(length=10000)


def exam_template_csv() -> str:
    output = StringIO()
    output.write("candidate_id,phase,exam_type,custom_exam_name,attempt,score,max_score,pass_percentage\n")
    output.write("candidate-id,PRE_ONBOARDING,FOUNDATION_EXAM,,1,78,100,60\n")
    output.write("candidate-id,POST_ONBOARDING,FINAL_EXAM,,1,82,100,60\n")
    output.write("candidate-id,PRE_ONBOARDING,CUSTOM_EXAM,Python Foundation,1,74,100,60\n")
    return output.getvalue()
