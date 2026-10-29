from io import BytesIO
import re
from typing import Any

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook

from app.core.enums import ExamPhase, ExamType
from app.db.postgresql import get_database
from app.schemas.common import utc_now
from app.utils.ids import new_id


def _normalize_exam_name(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _knowledge_check_column_name(exam_name: str) -> str | None:
    normalized = exam_name.strip().lower()
    compact = re.sub(r"[^a-z0-9]", "", normalized)
    if "python" in normalized or "python" in compact:
        return "knowledge_check_python_score"
    if ("java" in normalized or compact.startswith("java")) and "javascript" not in normalized and not compact.startswith("javascript"):
        return "knowledge_check_java_score"
    return None


def _extract_knowledge_exams(data: dict[str, Any]) -> list[dict[str, Any]]:
    coding_name = _normalize_exam_name(data.get("coding_name") or data.get("exam_name") or data.get("knowledge_check_name"))
    if coding_name and data.get("score") not in (None, ""):
        return [
            {
                "exam_name": coding_name,
                "score": float(data["score"]),
                "max_score": float(data.get("max_score") or 100),
                "pass_percentage": float(data.get("pass_percentage") or 0),
                "attempt": int(data.get("attempt") or 1),
            }
        ]

    ignore_fields = {
        "candidate_email",
        "email",
        "name",
        "college",
        "score",
        "max_score",
        "pass_percentage",
        "attempt",
        "percentage",
        "coding_name",
        "exam_name",
        "knowledge_check_name",
    }
    exams: list[dict[str, Any]] = []
    for key, value in data.items():
        if key in ignore_fields or value in (None, ""):
            continue
        try:
            score = float(value)
        except (TypeError, ValueError):
            continue
        exams.append(
            {
                "exam_name": key,
                "score": score,
                "max_score": 100.0,
                "pass_percentage": 0.0,
                "attempt": 1,
            }
        )
    return exams


async def upload_knowledge_check(file: UploadFile, user_id: str) -> dict[str, Any]:

    if not file.filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Only .xlsx supported")

    workbook = load_workbook(BytesIO(await file.read()))
    sheet = workbook.active

    headers = [
        str(cell.value).strip().lower() if cell.value else ""
        for cell in sheet[1]
    ]

    if not headers:
        raise HTTPException(status_code=400, detail="Empty Excel file")

    db = get_database()
    rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]

    processed = 0
    unmatched = []

    for idx, data in enumerate(rows, start=2):

        try:
            email = str(data.get("candidate_email") or data.get("email") or "").strip().lower()
            name = str(data.get("name") or "").strip()

            if not email and not name:
                raise ValueError("No identifier")

            candidate = None

            # -------------------------
            # MATCH CANDIDATE
            # -------------------------
            if email:
                candidate = await db.candidates.find_one({"email": email})

            if not candidate and name:
                possible = await db.candidates.find(
                    {"name": {"$regex": name, "$options": "i"}}
                ).to_list(length=5)

                if len(possible) == 1:
                    candidate = possible[0]

            if not candidate:
                unmatched.append({
                    "row": idx,
                    "email": email,
                    "name": name
                })
                continue

            exams = _extract_knowledge_exams(data)
            if not exams:
                unmatched.append({
                    "row": idx,
                    "email": email,
                    "name": name,
                    "error": "No knowledge check marks found in row",
                })
                continue

            candidate_updates: dict[str, Any] = {"updated_at": utc_now()}

            for exam in exams:
                score = exam["score"]
                attended = score > 0

                await db.execute(
                    """
                    INSERT INTO exam_results (
                        id, candidate_id, batch_id, phase, exam_type, exam_name,
                        attempt, score, max_score, pass_percentage, passed,
                        candidate_status_after, timeline_badge, created_by, created_at
                    )
                    VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,$12,$13,$14,$15)
                    ON CONFLICT (candidate_id, phase, exam_name, attempt)
                    DO UPDATE SET
                        score = EXCLUDED.score,
                        max_score = EXCLUDED.max_score,
                        passed = EXCLUDED.passed,
                        candidate_status_after = EXCLUDED.candidate_status_after,
                        created_by = EXCLUDED.created_by,
                        created_at = EXCLUDED.created_at
                    """,
                    new_id(),
                    candidate["id"],
                    candidate.get("batch_id"),
                    ExamPhase.POST_ONBOARDING.value,
                    ExamType.KNOWLEDGE_CHECK.value,
                    exam["exam_name"],
                    exam["attempt"],
                    score,
                    exam["max_score"],
                    exam["pass_percentage"],
                    attended,
                    candidate.get("status") or "ELIGIBLE",
                    "KNOWLEDGE_CHECK",
                    user_id,
                    utc_now()
                )
                column_name = _knowledge_check_column_name(exam["exam_name"])
                if column_name:
                    candidate_updates[column_name] = score
                processed += 1

            await db.update_candidate_row({"id": candidate["id"]}, candidate_updates)

        except Exception as e:
            unmatched.append({
                "row": idx,
                "error": str(e)
            })

    return {
        "processed": processed,
        "unmatched": unmatched[:50]
    }
